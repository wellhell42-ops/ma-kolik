"""Data export module - JSON, CSV, Excel."""

import csv
import json
import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


def export_json(data: Any, filepath: str, indent: int = 2) -> str:
    """Export data to JSON file.

    Args:
        data: Data to export (must be JSON serializable or have to_dict())
        filepath: Output file path
        indent: JSON indentation

    Returns:
        Absolute path to the created file
    """
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    serializable = _make_serializable(data)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(serializable, f, ensure_ascii=False, indent=indent)

    logger.info(f"Exported JSON: {path.absolute()}")
    return str(path.absolute())


def export_csv(data: list[dict], filepath: str) -> str:
    """Export list of dicts to CSV file.

    Args:
        data: List of dictionaries to export
        filepath: Output file path

    Returns:
        Absolute path to the created file
    """
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    if not data:
        logger.warning("No data to export to CSV")
        return str(path.absolute())

    rows = [_make_serializable(item) for item in data]
    if not rows:
        return str(path.absolute())

    # Collect all keys
    fieldnames = []
    for row in rows:
        if isinstance(row, dict):
            for key in row:
                if key not in fieldnames:
                    fieldnames.append(key)

    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            if isinstance(row, dict):
                writer.writerow(row)

    logger.info(f"Exported CSV: {path.absolute()}")
    return str(path.absolute())


def export_excel(data: dict[str, list], filepath: str) -> str:
    """Export data to Excel file with multiple sheets.

    Args:
        data: Dict mapping sheet names to lists of dicts
        filepath: Output file path (.xlsx)

    Returns:
        Absolute path to the created file
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        logger.error("openpyxl not installed. Run: pip install openpyxl")
        return ""

    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for sheet_name, rows in data.items():
        if not rows:
            continue

        ws = wb.create_sheet(title=sheet_name[:31])  # Excel 31 char limit
        serialized_rows = [_make_serializable(r) for r in rows]

        # Collect headers
        headers = []
        for row in serialized_rows:
            if isinstance(row, dict):
                for key in row:
                    if key not in headers:
                        headers.append(key)

        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=_turkish_header(header))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data
        for row_idx, row_data in enumerate(serialized_rows, 2):
            if not isinstance(row_data, dict):
                continue
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

        # Auto-fit column widths
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

        # Freeze header row
        ws.freeze_panes = "A2"

    if not wb.sheetnames:
        wb.create_sheet("Boş")

    wb.save(path)
    logger.info(f"Exported Excel: {path.absolute()}")
    return str(path.absolute())


def _make_serializable(obj: Any) -> Any:
    """Convert objects to JSON-serializable form."""
    if hasattr(obj, "to_dict"):
        return obj.to_dict()
    if isinstance(obj, list):
        return [_make_serializable(item) for item in obj]
    if isinstance(obj, dict):
        return {k: _make_serializable(v) for k, v in obj.items()}
    return obj


# Turkish translations for Excel headers
_HEADER_MAP = {
    "name": "İsim",
    "position": "Sıra",
    "played": "O",
    "won": "G",
    "drawn": "B",
    "lost": "M",
    "goals_for": "AG",
    "goals_against": "YG",
    "goal_difference": "AV",
    "points": "P",
    "rank": "Sıra",
    "team": "Takım",
    "value": "Değer",
    "matches_played": "Maç",
    "category": "Kategori",
    "home_team": "Ev Sahibi",
    "away_team": "Deplasman",
    "home_score": "Ev Skoru",
    "away_score": "Dep Skoru",
    "date": "Tarih",
    "time": "Saat",
    "status": "Durum",
    "week": "Hafta",
    "stadium": "Stadyum",
    "total_goals": "Toplam Gol",
    "goals_conceded": "Yenilen Gol",
    "clean_sheets": "Gol Yememe",
    "avg_goals_per_match": "Ort. Gol",
    "avg_conceded_per_match": "Ort. Yenilen",
    "minute": "Dakika",
    "league": "Lig",
}


def _turkish_header(key: str) -> str:
    """Convert field name to Turkish header."""
    return _HEADER_MAP.get(key, key.replace("_", " ").title())
